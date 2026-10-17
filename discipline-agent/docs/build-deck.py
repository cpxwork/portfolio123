
# -*- coding: utf-8 -*-
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

PRIMARY=RGBColor(0x67,0x50,0xA4); ON_PRIMARY=RGBColor(0xFF,0xFF,0xFF)
PRIM_CONT=RGBColor(0xEA,0xDD,0xFF); ON_PRIM_CONT=RGBColor(0x21,0x00,0x5D)
SECONDARY=RGBColor(0x62,0x5B,0x71); TERTIARY=RGBColor(0x7D,0x52,0x60)
TERT_CONT=RGBColor(0xFF,0xD8,0xE2); SURFACE=RGBColor(0xFF,0xFB,0xFE)
SURF_VAR=RGBColor(0xE7,0xE0,0xEC); SURF_CONT=RGBColor(0xF3,0xED,0xF7)
OUTLINE=RGBColor(0x79,0x74,0x7E); ON_SURFACE=RGBColor(0x1D,0x1B,0x20)
ON_SURF_VAR=RGBColor(0x49,0x45,0x4F); GREEN=RGBColor(0x2E,0x6B,0x4F)
GREEN_CONT=RGBColor(0xD7,0xEF,0xDF); GREEN_DARK=RGBColor(0x0F,0x3D,0x28)
FONT='微软雅黑'; MONO='Consolas'

prs=Presentation(); prs.slide_width=Emu(12192000); prs.slide_height=Emu(6858000)
SW,SH=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]

def solid(sh,c):
    sh.fill.solid(); sh.fill.fore_color.rgb=c; sh.line.fill.background(); sh.shadow.inherit=False

def card(sl,x,y,w,h,c=SURFACE,r=0.10,line=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h); solid(sh,c)
    try: sh.adjustments[0]=r
    except Exception: pass
    if line is not None:
        sh.line.color.rgb=line; sh.line.width=Pt(0.75)
    return sh

def bar(sl,x,y,w,h,c):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h); solid(sh,c)
    try: sh.adjustments[0]=0.5
    except Exception: pass
    return sh

def text(sl,x,y,w,h,runs,size=16,color=ON_SURFACE,bold=False,
         align=PP_ALIGN.LEFT,spacing=1.0,space_after=6):
    tb=sl.shapes.add_textbox(x,y,w,h); tf=tb.text_frame
    tf.word_wrap=True; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    if isinstance(runs,str): runs=[runs]
    for i,item in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=spacing; p.space_after=Pt(space_after)
        if isinstance(item,str): item=[(item,{})]
        for seg,st in item:
            r=p.add_run(); r.text=seg; f=r.font
            f.name=st.get('font',FONT); f.size=Pt(st.get('size',size))
            f.bold=st.get('bold',bold); f.color.rgb=st.get('color',color)
    return tb

def label(sl,x,y,t,c=PRIMARY,size=11):
    text(sl,x,y,Inches(5),Inches(0.28),[[(t,{'size':size,'bold':True,'color':c})]])

def header(sl,kicker,title,sub=None):
    bar(sl,Inches(0.62),Inches(0.52),Inches(0.42),Inches(0.055),PRIMARY)
    label(sl,Inches(0.62),Inches(0.66),kicker)
    text(sl,Inches(0.62),Inches(0.98),Inches(11.6),Inches(0.62),
         [[(title,{'size':29,'bold':True,'color':ON_SURFACE})]],size=29)
    if sub:
        text(sl,Inches(0.62),Inches(1.62),Inches(11.6),Inches(0.4),
             [[(sub,{'size':12.5,'color':ON_SURF_VAR})]],size=12.5)

def footer(sl,n):
    text(sl,Inches(11.35),Inches(6.94),Inches(1.4),Inches(0.3),
         [[('%02d'%n,{'size':11,'color':OUTLINE})]],align=PP_ALIGN.RIGHT)
    text(sl,Inches(0.62),Inches(6.94),Inches(7),Inches(0.3),
         [[('自律助手 · 项目介绍',{'size':10,'color':OUTLINE})]])

N=[0]
def nxt():
    N[0]+=1; return N[0]

def new():
    s=prs.slides.add_slide(BLANK); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=SURFACE; return s

# ============ 1 封面 ============
s=new()
card(s,Emu(0),Emu(0),SW,SH,SURFACE)
card(s,Inches(0),Inches(0),Inches(13.333),Inches(4.35),PRIM_CONT)
bar(s,Inches(0.9),Inches(0.95),Inches(0.55),Inches(0.07),PRIMARY)
text(s,Inches(0.9),Inches(1.18),Inches(11),Inches(0.4),
     [[('项目介绍',{'size':13,'bold':True,'color':ON_PRIM_CONT})]])
text(s,Inches(0.9),Inches(1.58),Inches(11.5),Inches(1.2),
     [[('自律助手',{'size':54,'bold':True,'color':ON_PRIM_CONT})]])
text(s,Inches(0.9),Inches(2.62),Inches(11.5),Inches(1.2),
     [[('从聊天里发现该做的事，按你现在的状态拆成小步，',{'size':18,'color':ON_PRIM_CONT})],
      [('不动就逐级升级 —— 全程可撤销、可审计。',{'size':18,'color':ON_PRIM_CONT})]])
for i,(k,v) in enumerate([('六级强制','L0 → L5'),('指令动作','7 种'),
                          ('测试状态','13 套全绿'),('还原保证','100% 可还原')]):
    x=Inches(0.9+i*2.95)
    card(s,x,Inches(4.8),Inches(2.7),Inches(1.25),SURF_CONT,0.16)
    text(s,x+Inches(0.26),Inches(5.03),Inches(2.2),Inches(0.35),
         [[(k,{'size':11,'color':ON_SURF_VAR})]])
    text(s,x+Inches(0.26),Inches(5.33),Inches(2.3),Inches(0.5),
         [[(v,{'size':20,'bold':True,'color':PRIMARY})]])
text(s,Inches(0.9),Inches(6.35),Inches(11),Inches(0.4),
     [[('陈鹏旭 · 产品与 AI 学习作品 · 2026-10-17',{'size':12,'color':ON_SURF_VAR})]])

# ============ 2 问题 ============
s=new(); footer(s,nxt())
header(s,'01 · 问题','该做的事，总是"知道"却不做',
       '待办躺在聊天记录里；提醒弹出即被划掉；没有下一步，于是不动。')
items=[('提醒会疲劳','角标、弹窗、提示音——同一招用三次就免疫，用户学会无视。'),
       ('不知道下一步','"写数学作业"太大，躺着的时候它不是一个能立刻开始的动作。'),
       ('强制没有刹车','真去动文件的做法，一旦出错就不可逆；用户不敢给权限。')]
for i,(t,d) in enumerate(items):
    x=Inches(0.62+i*4.06)
    card(s,x,Inches(2.35),Inches(3.78),Inches(2.2),SURF_CONT,0.13)
    bar(s,x+Inches(0.32),Inches(2.64),Inches(0.34),Inches(0.055),TERTIARY)
    text(s,x+Inches(0.32),Inches(2.9),Inches(3.2),Inches(0.4),
         [[(t,{'size':17,'bold':True,'color':ON_SURFACE})]])
    text(s,x+Inches(0.32),Inches(3.4),Inches(3.22),Inches(1.0),
         [[(d,{'size':12.5,'color':ON_SURF_VAR})]],spacing=1.25)
card(s,Inches(0.62),Inches(4.9),Inches(12.12),Inches(1.4),PRIM_CONT,0.20)
text(s,Inches(1.0),Inches(5.16),Inches(11.4),Inches(1.0),
     [[('我们要做的：',{'size':13,'bold':True,'color':ON_PRIM_CONT})],
      [('把"提醒"换成一条有起点、有下一步、失败会升级，但任何时刻都能一键收回的路径。',
        {'size':15,'color':ON_PRIM_CONT})]],spacing=1.3)

# ============ 3 闭环 ============
s=new(); footer(s,nxt())
header(s,'02 · 一句话说清','感知 → 判断 → 拆步 → 升级 → 执行 → 可还原',
       '六个环节串成一个闭环；用户完成任意一步，全部强制状态立即撤销。')
steps=[('感知','聊天记录\n时间 / 活跃度',PRIMARY),('判断','该不该学\n现在什么状态',PRIMARY),
       ('拆步','≤7 步 · 每步 ≤3 分钟\n从当前状态出发',SECONDARY),
       ('升级','L0 → L5\n每级留真实间隔',TERTIARY),
       ('执行','文件打乱 / 隔离\n先写 manifest',TERTIARY),
       ('还原','一键 revoke_all\n逐字节校验',GREEN)]
for i,(t,d,c) in enumerate(steps):
    x=Inches(0.62+i*2.06)
    card(s,x,Inches(2.4),Inches(1.85),Inches(2.45),SURF_CONT,0.15)
    bar(s,x+Inches(0.26),Inches(2.64),Inches(0.3),Inches(0.05),c)
    text(s,x+Inches(0.26),Inches(2.86),Inches(1.4),Inches(0.4),
         [[(t,{'size':17,'bold':True,'color':ON_SURFACE})]])
    text(s,x+Inches(0.26),Inches(3.38),Inches(1.62),Inches(1.2),
         [[(d,{'size':11,'color':ON_SURF_VAR})]],spacing=1.25)
    if i<5:
        text(s,x+Inches(1.86),Inches(3.35),Inches(0.24),Inches(0.4),
             [[('›',{'size':18,'color':OUTLINE})]])
card(s,Inches(0.62),Inches(5.15),Inches(12.12),Inches(1.15),GREEN_CONT,0.20)
text(s,Inches(1.0),Inches(5.4),Inches(11.4),Inches(0.7),
     [[('硬约束：任何一级用户完成该步 → 立即撤销全部强制状态；永不删除文件，只移动 / 重命名。',
        {'size':13.5,'bold':True,'color':GREEN_DARK})]])

# ============ 4 分级 ============
s=new(); footer(s,nxt())
header(s,'03 · 强制措施分级','L0 → L5：六级递进，越往后越难无视',
       '每级之间必须留足真实等待时间，禁止连跳；单日强制上限 4 次，达到后当天降级为纯提醒。')
rows=[('L0','notify','轻提醒：角标 / 提示音','前端渲染',PRIMARY),
      ('L1','popup_countdown','弹窗 + 倒计时 +「开始」按钮','前端渲染',PRIMARY),
      ('L2','spawn_bugs','满屏虫子，挡操作（带关闭倒计时）','前端渲染',SECONDARY),
      ('L3','screen_flash','整机闪屏，强制回到桌面','前端渲染',SECONDARY),
      ('L4','shuffle_desktop','打乱桌面文件（记 manifest）','A 侧执行',TERTIARY),
      ('L5','quarantine_files','移动指定文件到隔离区（需事前授权）','A 侧执行',TERTIARY)]
for i,(lv,act,desc,owner,c) in enumerate(rows):
    y=Inches(2.35+i*0.66)
    card(s,Inches(0.62),y,Inches(12.12),Inches(0.56),SURF_CONT if i%2==0 else SURFACE,0.28)
    bar(s,Inches(0.78),y+Inches(0.14),Inches(0.055),Inches(0.28),c)
    text(s,Inches(0.98),y+Inches(0.11),Inches(0.7),Inches(0.35),
         [[(lv,{'size':15,'bold':True,'color':c})]])
    text(s,Inches(1.72),y+Inches(0.13),Inches(2.7),Inches(0.35),
         [[(act,{'size':12,'bold':True,'color':ON_SURFACE,'font':MONO})]])
    text(s,Inches(4.5),y+Inches(0.13),Inches(5.4),Inches(0.35),
         [[(desc,{'size':12,'color':ON_SURF_VAR})]])
    text(s,Inches(10.3),y+Inches(0.13),Inches(2.25),Inches(0.35),
         [[(owner,{'size':11,'bold':True,'color':c})]],align=PP_ALIGN.RIGHT)
text(s,Inches(0.62),Inches(6.42),Inches(11.8),Inches(0.4),
     [[('L4 起才碰文件系统，且只在用户事前勾选授权 + 指定文件清单之后。',
        {'size':11.5,'color':ON_SURF_VAR})]])

# ============ 5 分工 ============
s=new(); footer(s,nxt())
header(s,'04 · 前后端分工','A 侧负责"判断"与"执行"，B 侧负责"表现"',
       '两侧只通过两条契约通信：指令 JSON 与状态模型。前端不参与任何文件操作。')
card(s,Inches(0.62),Inches(2.3),Inches(5.86),Inches(3.85),PRIM_CONT,0.10)
text(s,Inches(1.0),Inches(2.58),Inches(5.1),Inches(0.45),
     [[('A 侧 · 后端 + Agent + 本地层',{'size':17,'bold':True,'color':ON_PRIM_CONT})]])
for i,t in enumerate(['聊天待办提取（只用脱敏摘要）','摄像头 / 时间 → 状态估计',
                      '状态驱动拆小步（≤7 步）','升级决策：模型 + 物理闸门',
                      '文件执行：打乱 / 隔离 / 还原','审计：读了什么、依据什么、做了什么']):
    text(s,Inches(1.0),Inches(3.16+i*0.47),Inches(5.2),Inches(0.4),
         [[('•  ',{'size':12,'color':PRIMARY}),(t,{'size':12.5,'color':ON_PRIM_CONT})]])
card(s,Inches(6.88),Inches(2.3),Inches(5.86),Inches(3.85),TERT_CONT,0.10)
text(s,Inches(7.26),Inches(2.58),Inches(5.1),Inches(0.45),
     [[('B 侧 · 前端表现',{'size':17,'bold':True,'color':RGBColor(0x31,0x1D,0x2A)})]])
for i,t in enumerate(['提醒 UI、弹窗与倒计时','满屏虫子动画','整机闪屏（L3）',
                      '任务列表与小步进度','授权与「指定文件」设置页','可撤销入口、状态可视化']):
    text(s,Inches(7.26),Inches(3.16+i*0.47),Inches(5.2),Inches(0.4),
         [[('•  ',{'size':12,'color':TERTIARY}),(t,{'size':12.5,'color':RGBColor(0x31,0x1D,0x2A)})]])

# ============ 6 指令 JSON ============
s=new(); footer(s,nxt())
header(s,'05 · 契约一','指令 JSON —— 唯一的下发格式','字段与枚举先定死，不许擅改；变更必须显式记录。')
code=[('{',ON_SURF_VAR),
      ('  "id": "ins_01",',ON_SURFACE),
      ('  "level": "L2",',PRIMARY),
      ('  "action": "spawn_bugs",',PRIMARY),
      ('  "taskId": "task_math_01",',ON_SURFACE),
      ('  "step": { "index": 1, "label": "下床", "stateAtPlan": "躺床" },',ON_SURFACE),
      ('  "issuedAt": "2026-10-16T20:30:00+08:00",',ON_SURF_VAR),
      ('  "expiresAt": "2026-10-16T20:35:00+08:00",',ON_SURF_VAR),
      ('  "reversible": true,',GREEN),
      ('  "manifestId": null',ON_SURF_VAR),
      ('}',ON_SURF_VAR)]
card(s,Inches(0.62),Inches(2.28),Inches(7.3),Inches(3.95),SURF_CONT,0.08)
tb=text(s,Inches(0.95),Inches(2.52),Inches(6.7),Inches(3.5),[])
tf=tb.text_frame
for i,(line,c) in enumerate(code):
    p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.line_spacing=1.18; p.space_after=Pt(0)
    r=p.add_run(); r.text=line
    r.font.name=MONO; r.font.size=Pt(11.5); r.font.color.rgb=c
text(s,Inches(8.22),Inches(2.28),Inches(4.5),Inches(0.4),
     [[('要点',{'size':15,'bold':True,'color':ON_SURFACE})]])
pts=[('7 种动作','notify / popup_countdown / spawn_bugs / screen_flash / shuffle_desktop / quarantine_files / revoke_all'),
     ('可逆优先','reversible 标记 + manifestId 回链，任何时刻能还原'),
     ('payload 扩展','quarantine_files → targets[]；shuffle_desktop → targetDir'),
     ('校验必做','id / level / action / 时间戳 / step.index；不合法抛错并附 errors')]
for i,(t,d) in enumerate(pts):
    y=Inches(2.68+i*0.9)
    card(s,Inches(8.22),y,Inches(4.5),Inches(0.78),SURF_CONT,0.18)
    text(s,Inches(8.46),y+Inches(0.08),Inches(4.0),Inches(0.3),
         [[(t,{'size':12.5,'bold':True,'color':PRIMARY})]])
    text(s,Inches(8.46),y+Inches(0.35),Inches(4.05),Inches(0.4),
         [[(d,{'size':9.5,'color':ON_SURF_VAR})]],spacing=1.1)

# ============ 7 状态模型 ============
s=new(); footer(s,nxt())
header(s,'06 · 契约二','状态模型 —— A → B 的唯一输入',
       '前端靠 explain / evidence 把冷冰冰的指令还原成人话。')
card(s,Inches(0.62),Inches(2.28),Inches(5.86),Inches(2.05),SURF_CONT,0.10)
text(s,Inches(0.95),Inches(2.52),Inches(5.2),Inches(0.4),
     [[('六个状态',{'size':15,'bold':True,'color':PRIMARY})]])
for i,t in enumerate(['lying_bed · 躺床','seated_desk · 在桌前','standing_or_moving · 走动中',
                      'away · 离开','learning · 学习中','unknown · 未知']):
    text(s,Inches(0.95+(i%2)*2.85),Inches(2.96+(i//2)*0.42),Inches(2.85),Inches(0.36),
         [[('• ',{'size':11,'color':OUTLINE}),(t,{'size':12,'color':ON_SURFACE})]])
card(s,Inches(6.88),Inches(2.28),Inches(5.86),Inches(2.05),SURF_CONT,0.10)
text(s,Inches(7.21),Inches(2.52),Inches(5.2),Inches(0.4),
     [[('必须可解释',{'size':15,'bold':True,'color':PRIMARY})]])
text(s,Inches(7.21),Inches(2.98),Inches(5.25),Inches(1.2),
     [[('"因为摄像头判定为床区域卧姿，且已 18 分钟无操作。"',{'size':12.5,'color':ON_SURFACE})],
      [('evidence 带证据链，explain 给成人话。',{'size':11.5,'color':ON_SURF_VAR})]],spacing=1.3)
card(s,Inches(0.62),Inches(4.55),Inches(12.12),Inches(1.75),GREEN_CONT,0.14)
text(s,Inches(1.0),Inches(4.8),Inches(11.5),Inches(1.35),
     [[('三条安全规则',{'size':14,'bold':True,'color':GREEN_DARK})],
      [('① unknown 是合法且常用的状态 —— 宁可 unknown 也不瞎猜；低置信度不拆激进小步、不进强制升级。',
        {'size':12,'color':GREEN_DARK})],
      [('② 摄像头证据不能单独触发 L2+ —— 必须叠加时间规则或用户确认。',{'size':12,'color':GREEN_DARK})],
      [('③ 用户手动纠正（source=manual）权重最高，直接覆盖自动判定。',{'size':12,'color':GREEN_DARK})]],spacing=1.22)

# ============ 8 物理闸门 ============
s=new(); footer(s,nxt())
header(s,'07 · 设计红线','模型可以决策，但数字只能算出来',
       '强制措施必须可预测、可解释 —— 所以把"决定权"和"执行权"分开。')
gates=[('闸门 1','完成信号优先','任何一步完成 → 无条件撤销全部强制，模型说了不算。'),
       ('闸门 2','时间门禁','距上次升级不足该级间隔，不许升；禁止连跳。'),
       ('闸门 3','单日上限 + 授权','单日强制上限 4 次；L5 隔离需事前授权，未授权封顶 L4。'),
       ('闸门 4','动作不由模型选','action 由 level 映射表决定，模型不能自选动作。'),
       ('闸门 5','清单不由模型给','隔离目标只来自用户事前指定的文件清单。')]
for i,(k,t,d) in enumerate(gates):
    y=Inches(2.3+i*0.82)
    card(s,Inches(0.62),y,Inches(8.4),Inches(0.7),SURF_CONT,0.22)
    bar(s,Inches(0.8),y+Inches(0.2),Inches(0.05),Inches(0.3),PRIMARY)
    text(s,Inches(1.02),y+Inches(0.18),Inches(0.9),Inches(0.35),
         [[(k,{'size':11,'bold':True,'color':PRIMARY})]])
    text(s,Inches(1.95),y+Inches(0.18),Inches(2.1),Inches(0.35),
         [[(t,{'size':13,'bold':True,'color':ON_SURFACE})]])
    text(s,Inches(4.1),y+Inches(0.19),Inches(4.75),Inches(0.35),
         [[(d,{'size':11,'color':ON_SURF_VAR})]])
card(s,Inches(9.32),Inches(2.3),Inches(3.42),Inches(3.98),PRIM_CONT,0.10)
text(s,Inches(9.62),Inches(2.56),Inches(2.95),Inches(3.5),
     [[('为什么这样分',{'size':14,'bold':True,'color':ON_PRIM_CONT})],
      [('',{'size':6})],
      [('LLM 负责"怎么劝"——语气、时机、拆步；',{'size':12,'color':ON_PRIM_CONT})],
      [('',{'size':4})],
      [('确定性代码负责"能不能动"——数字、阈值、文件操作。',{'size':12,'color':ON_PRIM_CONT})],
      [('',{'size':6})],
      [('这样任何一次强制行为都能复盘到具体规则，而不是"模型当时觉得"。',
        {'size':11.5,'color':ON_PRIM_CONT})]],spacing=1.25)

# ============ 9 可还原 ============
s=new(); footer(s,nxt())
header(s,'08 · 安全底座','永不删除，只移动；任何操作都能逐字节还原',
       '用户的信任来自"我随时能收回来"。')
flow=[('写 manifest','操作前先把\n原路径 ↔ 新路径\n写进账本',PRIMARY),
      ('执行','打乱 = 重命名\n隔离 = 移动到\n隔离区（加密）',SECONDARY),
      ('逐条回写','每完成一条\n就回写账本\n崩溃也能还原',SECONDARY),
      ('revoke_all','一键撤销\n恢复原位\n校验 sha256',GREEN)]
for i,(t,d,c) in enumerate(flow):
    x=Inches(0.62+i*3.13)
    card(s,x,Inches(2.35),Inches(2.85),Inches(2.25),SURF_CONT,0.13)
    bar(s,x+Inches(0.3),Inches(2.58),Inches(0.32),Inches(0.055),c)
    text(s,x+Inches(0.3),Inches(2.82),Inches(2.3),Inches(0.4),
         [[(t,{'size':16,'bold':True,'color':ON_SURFACE})]])
    text(s,x+Inches(0.3),Inches(3.32),Inches(2.4),Inches(1.1),
         [[(d,{'size':11.5,'color':ON_SURF_VAR})]],spacing=1.22)
    if i<3:
        text(s,x+Inches(2.86),Inches(3.28),Inches(0.28),Inches(0.4),
             [[('›',{'size':18,'color':OUTLINE})]])
card(s,Inches(0.62),Inches(4.9),Inches(12.12),Inches(1.4),SURF_CONT,0.18)
text(s,Inches(1.0),Inches(5.14),Inches(11.5),Inches(1.0),
     [[('其他约束：',{'size':13,'bold':True,'color':PRIMARY})],
      [('只操作授权目录之内（拒绝系统目录 / 盘符根 / 用户主目录本身）；restore 幂等，原位被占记 conflict 不覆盖；跳过符号链接与点开头文件。',
        {'size':12,'color':ON_SURF_VAR})]],spacing=1.3)

# ============ 10 组成模块 ============
s=new(); footer(s,nxt())
header(s,'09 · 系统组成','六个环节，各自独立可测',
       '每个模块都有对应的测试套件；改一处不会悄悄弄坏别处。')
mods=[('感知','chat-sense / state-sensor','聊天待办提取、摄像头 + 时间规则状态估计'),
      ('规划','task-planner','状态驱动拆小步 + 状态变化重排'),
      ('决策','escalation + llm/roles','模型决策，物理闸门兜底'),
      ('执行','executor / manifest / instructions','文件操作、可逆账本、指令总线与审计'),
      ('学习','learning-kernel + learning/store','间隔重复调度内核与学习记录存储'),
      ('审计','audit','审计查询、用户纠正、误判申诉')]
for i,(k,f,d) in enumerate(mods):
    x=Inches(0.62+(i%3)*4.06); y=Inches(2.28+(i//3)*1.95)
    card(s,x,y,Inches(3.78),Inches(1.7),SURF_CONT,0.14)
    bar(s,x+Inches(0.3),y+Inches(0.24),Inches(0.3),Inches(0.05),PRIMARY)
    text(s,x+Inches(0.3),y+Inches(0.44),Inches(3.2),Inches(0.32),
         [[(k,{'size':15,'bold':True,'color':ON_SURFACE})]])
    text(s,x+Inches(0.3),y+Inches(0.8),Inches(3.28),Inches(0.3),
         [[(f,{'size':9.5,'bold':True,'color':PRIMARY,'font':MONO})]])
    text(s,x+Inches(0.3),y+Inches(1.1),Inches(3.28),Inches(0.5),
         [[(d,{'size':11,'color':ON_SURF_VAR})]],spacing=1.15)
text(s,Inches(0.62),Inches(6.3),Inches(11.8),Inches(0.35),
     [[('测试：13 套 / 13 通过 / 0 失败（node run-tests.js，零第三方依赖）。',
        {'size':12,'bold':True,'color':GREEN})]])

# ============ 11 现状 ============
s=new(); footer(s,nxt())
header(s,'10 · 现在的状态','已经能跑通的部分','以下是本仓库实际可复跑的验证结果，不含推测。')
done=[('指令通道 + 可逆账本','下发 / 确认 / 完成 / 撤销全链路，sha256 逐字节校验还原'),
      ('六级升级状态机','L0→L5 逐级触发、时间门禁、单日上限、授权封顶、审计回链'),
      ('状态估计','四类场景判定 + 低置信度降级 + 摄像头证据不单独触发 L2+'),
      ('任务拆步','≤7 步、每步 ≤3 分钟、第一步从当前状态出发、状态变化重排'),
      ('本地执行','打乱 / 隔离 / 还原幂等，冲突不覆盖，崩溃可还原'),
      ('学习子系统','间隔重复调度内核 + 学习记录存储（跨进程 JSON 边界调用）')]
for i,(t,d) in enumerate(done):
    y=Inches(2.26+i*0.7)
    card(s,Inches(0.62),y,Inches(12.12),Inches(0.58),SURF_CONT if i%2==0 else SURFACE,0.26)
    bar(s,Inches(0.82),y+Inches(0.15),Inches(0.05),Inches(0.28),GREEN)
    text(s,Inches(1.05),y+Inches(0.13),Inches(3.0),Inches(0.34),
         [[(t,{'size':12.5,'bold':True,'color':ON_SURFACE})]])
    text(s,Inches(4.2),y+Inches(0.145),Inches(8.3),Inches(0.34),
         [[(d,{'size':11,'color':ON_SURF_VAR})]])
card(s,Inches(0.62),Inches(6.42),Inches(12.12),Inches(0.52),PRIM_CONT,0.3)
text(s,Inches(1.0),Inches(6.53),Inches(11.4),Inches(0.4),
     [[('接入真实摄像头后端与真实模型 API 后，即可端到端试用。',
        {'size':12,'bold':True,'color':ON_PRIM_CONT})]])

# ============ 12 结语 ============
s=new()
card(s,Emu(0),Emu(0),SW,SH,PRIM_CONT)
bar(s,Inches(5.95),Inches(1.5),Inches(1.45),Inches(0.07),PRIMARY)
text(s,Inches(1.2),Inches(1.9),Inches(10.9),Inches(1.0),
     [[('让"知道该做"变成"已经在做"',{'size':34,'bold':True,'color':ON_PRIM_CONT})]],
     align=PP_ALIGN.CENTER)
text(s,Inches(1.9),Inches(3.05),Inches(9.5),Inches(1.6),
     [[('判断交给模型，动文件交给规则；',{'size':18,'color':ON_PRIM_CONT})],
      [('每一步都留退路，每一级都讲得出理由。',{'size':18,'color':ON_PRIM_CONT})]],
     align=PP_ALIGN.CENTER,spacing=1.45)
card(s,Inches(4.05),Inches(4.8),Inches(5.2),Inches(0.95),SURFACE,0.22)
text(s,Inches(4.3),Inches(5.03),Inches(4.7),Inches(0.5),
     [[('13 套测试全绿 · 100% 可还原 · 全程可审计',{'size':13.5,'bold':True,'color':PRIMARY})]],
     align=PP_ALIGN.CENTER)
text(s,Inches(1.2),Inches(6.05),Inches(10.9),Inches(0.4),
     [[('陈鹏旭 · 产品与 AI 学习作品',{'size':12,'color':ON_PRIM_CONT})]],align=PP_ALIGN.CENTER)

out='docs/自律助手-项目介绍.pptx'
prs.save(out)
print('SAVED', out, 'slides =', len(prs.slides._sldIdLst))
